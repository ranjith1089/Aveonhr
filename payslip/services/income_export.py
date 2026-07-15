"""Export the income ledger as an .xlsx matching the original sheet layout."""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from django.db.models import DecimalField, Sum, Value
from django.db.models.functions import Coalesce

from ..models import ClientBilling

HEADERS = [
    "Client", "Academic Year", "Student Count", "Rate", "One Time Payment",
    "Taxable Value", "GST 18%", "Net Amount", "Previous Pending", "Received",
    "Balance", "Implementation Engineer", "Invoice Status", "Remarks",
    "Agreement Status",
]
MONEY_COLS = {4, 5, 6, 7, 8, 9, 10, 11}  # 1-based column numbers with money format


def build_income_workbook(org) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    billings = (
        ClientBilling.objects.filter(client__organization=org)
        .select_related("client")
        .annotate(received_sum=Coalesce(
            Sum("payments__amount"), Value(Decimal("0")),
            output_field=DecimalField(max_digits=14, decimal_places=2),
        ))
        .order_by("client__name", "year_start")
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Income"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    fy_totals: dict[str, Decimal] = {}
    for b in billings:
        balance = b.total_due - b.received_sum
        fy_totals[b.academic_year] = fy_totals.get(b.academic_year, Decimal("0")) + balance
        ws.append([
            b.client.name,           # repeated per row (filterable, unlike the original)
            b.academic_year,
            b.student_count,
            b.rate,
            b.one_time_payment,
            b.taxable_value,
            b.gst_amount,
            b.net_amount,
            b.previous_pending,
            b.received_sum,
            balance,
            b.engineer,
            b.get_invoice_status_display() if b.invoice_status else "",
            b.remarks,
            b.client.agreement_status,
        ])

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in MONEY_COLS and cell.value is not None:
                cell.number_format = "#,##0.00"

    # Footer: per-FY outstanding + grand total, like the original sheet.
    ws.append([])
    grand = Decimal("0")
    for year in sorted(fy_totals):
        ws.append(["", "", "", "", "", "", "", "", "", year, fy_totals[year]])
        ws.cell(row=ws.max_row, column=11).number_format = "#,##0.00"
        grand += fy_totals[year]
    ws.append(["", "", "", "", "", "", "", "", "", "TOTAL :", grand])
    ws.cell(row=ws.max_row, column=11).number_format = "#,##0.00"
    ws.cell(row=ws.max_row, column=10).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=11).font = Font(bold=True)

    widths = [26, 13, 13, 10, 15, 14, 13, 14, 15, 14, 14, 20, 24, 40, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = BytesIO()
    wb.save(buf)  # in-memory only - Vercel has no writable disk
    return buf.getvalue()

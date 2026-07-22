"""Export a payroll run's register as an .xlsx, matching the source
sheet's original column layout."""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

HEADERS = [
    "S.No", "Name", "Designation", "Total Working Days", "CL Credit",
    "Emp Leave Days", "LOP Days", "Present Days", "Pay Days", "Basic", "DA",
    "HRA", "Transport Allowance", "Food Allowance", "Internet Allowance",
    "Salary Arrear / Allowance", "Gross Salary", "ESI Employee", "ESI Employer",
    "PF Employee", "PF Employer", "Salary Advance", "TDS", "Total Deductions",
    "Net Payable", "Remarks",
]
MONEY_COLS = set(range(10, 26))  # Basic..Net Payable (1-based columns)
HEADER_ROW = 3                   # row 1 = title, row 2 = spacer


def build_payroll_register_workbook(org, period) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    from ..models import PayslipEntry

    entries = (
        PayslipEntry.objects.filter(organization=org, run__period=period)
        .select_related("employee")
        .order_by("employee__name")
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll Register"

    # Title banner across the full table width, matching the printed
    # salary statement's heading.
    last_col = get_column_letter(len(HEADERS))
    ws["A1"] = f"Salary Statement For The Month Of {period:%B %Y}"
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.append([])                    # row 2 - spacer
    ws.append(HEADERS)               # row 3 - column headers
    for cell in ws[HEADER_ROW]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    totals = {"gross": Decimal("0"), "deductions": Decimal("0"), "net": Decimal("0")}
    for serial, e in enumerate(entries, start=1):
        totals["gross"] += e.gross_salary
        totals["deductions"] += e.total_deductions
        totals["net"] += e.net_payable
        ws.append([
            serial, e.employee.name, e.employee.designation, e.total_working_days,
            e.cl_credit, e.emp_leave_days, e.lop_days, e.present_days, e.pay_days,
            e.basic, e.da, e.hra, e.transport_allowance, e.food_allowance,
            e.internet_allowance, e.salary_arrear_allowance, e.gross_salary,
            e.esi_employee, e.esi_employer, e.pf_employee, e.pf_employer,
            e.salary_advance, e.tds, e.total_deductions, e.net_payable, e.remarks,
        ])

    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        for cell in row:
            if cell.column in MONEY_COLS and cell.value is not None:
                cell.number_format = "#,##0.00"

    ws.append([])
    ws.append(["", "", "", "", "", "", "", "", "", "", "", "", "", "", "",
               "TOTAL:", totals["gross"], "", "", "", "", "", "",
               totals["deductions"], totals["net"]])
    for col in (17, 24, 25):
        ws.cell(row=ws.max_row, column=col).number_format = "#,##0.00"
        ws.cell(row=ws.max_row, column=col).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=16).font = Font(bold=True)

    widths = [6, 24, 18, 10, 9, 11, 9, 10, 9, 11, 10, 10, 12, 11, 11, 14, 13,
             12, 12, 11, 11, 12, 10, 13, 13, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_employee_workbook(org, field_names: list[str], is_active=None) -> bytes:
    """Build an employee master export workbook with selected fields.

    Args:
        org: Organization to export for
        field_names: List of field names to include (e.g., ['name', 'designation', ...])
        is_active: None (all), True (active only), or False (inactive only)

    Returns: xlsx bytes
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    from ..models import Employee

    employees = Employee.objects.filter(organization=org)
    if is_active is not None:
        employees = employees.filter(is_active=is_active)
    employees = employees.order_by("name")

    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"

    # Mapping of field names to display labels and value formatters
    field_info = {
        "employee_code": ("Employee Code", lambda e: e.employee_code),
        "name": ("Name", lambda e: e.name),
        "designation": ("Designation", lambda e: e.designation or ""),
        "department": ("Department", lambda e: e.department or ""),
        "doj": ("Date of Joining", lambda e: (e.doj.strftime("%d-%m-%Y") if e.doj else "")),
        "relieving_date": ("Relieving Date", lambda e: (e.relieving_date.strftime("%d-%m-%Y") if e.relieving_date else "")),
        "is_active": ("Active", lambda e: ("Yes" if e.is_active else "No")),
        "employment_status": ("Employment Status", lambda e: e.get_employment_status_display()),
        "current_monthly_package": ("Monthly Package (₹)", lambda e: float(e.current_monthly_package)),
        "is_esi_eligible": ("ESI Eligible", lambda e: ("Yes" if e.is_esi_eligible else "No")),
        "is_pf_applicable": ("PF Applicable", lambda e: ("Yes" if e.is_pf_applicable else "No")),
        "pan_number": ("PAN Number", lambda e: e.pan_number or ""),
        "pf_number": ("PF Number", lambda e: e.pf_number or ""),
        "pf_uan": ("PF UAN", lambda e: e.pf_uan or ""),
        "esi_number": ("ESI Number", lambda e: e.esi_number or ""),
        "date_of_birth": ("Date of Birth", lambda e: (e.date_of_birth.strftime("%d-%m-%Y") if e.date_of_birth else "")),
        "blood_group": ("Blood Group", lambda e: e.blood_group or ""),
        "marital_status": ("Marital Status", lambda e: e.marital_status or ""),
        "aadhar_no": ("Aadhaar Number", lambda e: e.aadhar_no or ""),
        "address": ("Address", lambda e: e.address or ""),
        "personal_email": ("Personal Email", lambda e: e.personal_email or ""),
        "official_email": ("Official Email", lambda e: e.official_email or ""),
        "contact_no": ("Contact Number", lambda e: e.contact_no or ""),
        "official_no": ("Official Number", lambda e: e.official_no or ""),
        "emergency_no": ("Emergency Number", lambda e: e.emergency_no or ""),
        "agreement_signed": ("Agreement Signed", lambda e: ("Yes" if e.agreement_signed else "No")),
        "agreement_sign_date": ("Agreement Sign Date", lambda e: (e.agreement_sign_date.strftime("%d-%m-%Y") if e.agreement_sign_date else "")),
        "biometric_id": ("Biometric ID", lambda e: e.biometric_id or ""),
        "reason_for_leaving": ("Reason for Leaving", lambda e: e.reason_for_leaving or ""),
        "bank_name": ("Bank Name", lambda e: e.bank_name or ""),
        "bank_account_number": ("Account Number", lambda e: e.bank_account_number or ""),
        "ifsc_code": ("IFSC Code", lambda e: e.ifsc_code or ""),
    }

    # Build header row
    headers = [field_info[fn][0] for fn in field_names if fn in field_info]
    ws.append(headers)

    # Style header row
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Money columns (for conditional formatting)
    money_cols = set()
    for i, field_name in enumerate(field_names, start=1):
        if field_name == "current_monthly_package":
            money_cols.add(i)

    # Add employee data rows
    for employee in employees:
        row = []
        for field_name in field_names:
            if field_name in field_info:
                formatter = field_info[field_name][1]
                row.append(formatter(employee))
        ws.append(row)

    # Format money columns
    for row_idx, _ in enumerate(employees, start=2):
        for col_idx in money_cols:
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.number_format = "#,##0.00"
            cell.alignment = Alignment(horizontal="right")

    # Set column widths
    for col_idx, field_name in enumerate(field_names, start=1):
        # Adaptive width based on field type
        if field_name in ("address", "reason_for_leaving"):
            width = 35
        elif field_name in ("name", "designation", "department"):
            width = 25
        else:
            width = 15
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

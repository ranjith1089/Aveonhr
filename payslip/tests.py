from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse


def make_member(username, *, admin=False, org=None, **rights):
    """Create a user with an explicit membership.

    rights: short module keys (income=True, payroll=True, ...) - anything
    not named defaults to False so tests stay precise. Returns the user;
    reach the org via user.membership.organization.
    """
    from payslip.models import Membership, Organization
    user = User.objects.create_user(username, f"{username}@x.com", "pass12345")
    if org is None:
        org = Organization.objects.create(company_name=f"{username} org")
    field_values = {
        field: bool(rights.get(key, False))
        for key, field in Membership.MODULE_FIELDS.items()
    }
    Membership.objects.create(
        user=user, organization=org,
        role=Membership.Role.ADMIN if admin else Membership.Role.MEMBER,
        **field_values,
    )
    return user


# Valid payload for the CURRENT proposal form (bundle mode, DD/MM/YYYY date).
PROPOSAL_PAYLOAD = {
    "to_address": "The Principal",
    "client_name": "ABC College of Arts and Science",
    "client_address": "Coimbatore, Tamil Nadu",
    "proposal_date": "13/02/2026",
    "prepared_by": "Aveon Infotech Private Limited",
    "selection_mode": "BUNDLE",
    "bundle": "CMS_FULL",
    "pricing_model": "PER_STUDENT",
    "price_per_unit": "850",
    "minimum_student_commitment": "1000",
    "one_time_implementation_fee": "350000",
    "gst_percent": "18",
    "authorized_signatory_name": "Parvathi G",
    "authorized_signatory_designation": "Chief Executive Officer",
}


class ProposalQuotationViewTests(TestCase):
    def setUp(self):
        # create_user auto-provisions an admin membership on first request.
        self.user = User.objects.create_user("tester", "tester@example.com", "pass12345")
        self.client.force_login(self.user)

    def test_get_proposal_quotation_page(self):
        response = self.client.get(reverse("proposal_quotation"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Proposal Builder")

    def test_generate_creates_preview_download_and_history(self):
        from decimal import Decimal
        from payslip.models import ProposalRecord
        response = self.client.post(reverse("proposal_quotation"), PROPOSAL_PAYLOAD)
        self.assertEqual(response.status_code, 200)
        self.assertIn("preview_url", response.context)
        self.assertIn("download_url", response.context)

        record = ProposalRecord.objects.get()
        self.assertEqual(record.client_name, "ABC College of Arts and Science")
        self.assertEqual(record.revision, 1)
        self.assertEqual(record.created_by, self.user)
        # (850 * 1000 + 350000) * 1.18 = 14,16,000
        self.assertEqual(record.total_amount, Decimal("1416000.00"))
        self.assertIn("ABC College of Arts and Science", record.html)
        self.assertEqual(record.form_data["bundle"], "CMS_FULL")
        self.assertNotIn("client_logo", record.form_data)

    def test_preview_serves_generated_html(self):
        response = self.client.post(reverse("proposal_quotation"), PROPOSAL_PAYLOAD)
        preview = self.client.get(response.context["preview_url"])
        self.assertEqual(preview.status_code, 200)
        self.assertIn(b"ABC College of Arts and Science", preview.content)

    def test_bundle_mode_requires_bundle(self):
        from payslip.models import ProposalRecord
        payload = {**PROPOSAL_PAYLOAD, "bundle": ""}
        response = self.client.post(reverse("proposal_quotation"), payload)
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.context["form"].errors)
        self.assertEqual(ProposalRecord.objects.count(), 0)


class IncomeModuleTests(TestCase):
    def setUp(self):
        from .models import IncomeClient
        self.staff = make_member("inc_staff", income=True, implementation=True)
        self.org = self.staff.membership.organization
        # Member of the same org WITHOUT the income right -> 403.
        self.plain = make_member("inc_plain", org=self.org)
        self.client_obj = IncomeClient.objects.create(name="Test College",
                                                      organization=self.org)

    def test_gst_math_and_quantization(self):
        from decimal import Decimal
        from .models import ClientBilling
        b = ClientBilling(client=self.client_obj, academic_year="2025-2026",
                          student_count=1234, rate=Decimal("250"))
        b.save()
        self.assertEqual(b.taxable_value, Decimal("308500.00"))
        self.assertEqual(b.gst_amount, Decimal("55530.00"))
        self.assertEqual(b.net_amount, Decimal("364030.00"))
        self.assertEqual(b.year_start, 2025)

    def test_override_amounts_untouched(self):
        from decimal import Decimal
        from .models import ClientBilling
        b = ClientBilling(client=self.client_obj, academic_year="2024-2025",
                          override_amounts=True, net_amount=Decimal("850000.02"))
        b.save()
        self.assertEqual(b.net_amount, Decimal("850000.02"))
        self.assertEqual(b.gst_amount, Decimal("0"))

    def test_balance_with_part_payments_and_overpayment(self):
        from decimal import Decimal
        from .models import ClientBilling, PaymentReceipt
        b = ClientBilling(client=self.client_obj, academic_year="2025-2026",
                          override_amounts=True, net_amount=Decimal("100000"),
                          previous_pending=Decimal("20000"))
        b.save()
        PaymentReceipt.objects.create(billing=b, amount=Decimal("50000"))
        PaymentReceipt.objects.create(billing=b, amount=Decimal("80000"))
        self.assertEqual(b.received_total, Decimal("130000"))
        self.assertEqual(b.balance, Decimal("-10000"))  # overpaid -> advance

    def test_academic_year_normalization(self):
        from .forms_income import ClientBillingForm
        form = ClientBillingForm(data={"academic_year": "2025-26",
                                       "override_amounts": True, "net_amount": "1000",
                                       "one_time_payment": "0", "gst_amount": "0",
                                       "previous_pending": "0"})
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data["academic_year"], "2025-2026")

    def test_to_decimal_handles_currency_and_text(self):
        from decimal import Decimal
        from .services.income_import import _to_decimal
        self.assertEqual(_to_decimal("₹1,23,456.75"), Decimal("123456.75"))
        self.assertEqual(_to_decimal(1500), Decimal("1500.00"))
        self.assertIsNone(_to_decimal("Around 3 Lakhs"))
        self.assertIsNone(_to_decimal(""))

    def test_access_control(self):
        # anonymous -> login redirect
        r = self.client.get("/income/")
        self.assertEqual(r.status_code, 302)
        self.assertIn("/accounts/login/", r["Location"])
        # member without the income right -> friendly 403
        self.client.force_login(self.plain)
        resp = self.client.get("/income/")
        self.assertEqual(resp.status_code, 403)
        self.assertContains(resp, "organization admin", status_code=403)
        # member with the right -> 200
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get("/income/").status_code, 200)

    def test_seed_totals_match_sheet_footer(self):
        from decimal import Decimal
        from payslip.management.commands.seed_income import _rows
        rows = _rows().rows
        total = sum((r.net_amount + r.previous_pending - r.received for r in rows), Decimal("0"))
        self.assertEqual(total, Decimal("6281796.42"))


class IncomeAnalyticsTests(TestCase):
    """Forecast and analytics math for the Income module."""

    def setUp(self):
        from decimal import Decimal
        from payslip.models import ClientBilling, IncomeClient, PaymentReceipt
        self.staff = make_member("inc_staff2", income=True)
        self.org = self.staff.membership.organization
        self.client.force_login(self.staff)

        # Rate-based client with two consecutive years: 1000 -> 1100 students (10% growth)
        grower = IncomeClient.objects.create(name="Growth College", organization=self.org)
        ClientBilling.objects.create(client=grower, academic_year="2024-2025",
                                     student_count=1000, rate=Decimal("100"))
        b2 = ClientBilling.objects.create(client=grower, academic_year="2025-2026",
                                          student_count=1100, rate=Decimal("100"))
        PaymentReceipt.objects.create(billing=b2, amount=Decimal("50000"))

        # Fixed-fee client
        fixed = IncomeClient.objects.create(name="Fixed School", organization=self.org)
        ClientBilling.objects.create(client=fixed, academic_year="2025-2026",
                                     override_amounts=True, net_amount=Decimal("350000"))

        # Discontinued client - must be excluded from forecast
        gone = IncomeClient.objects.create(name="Gone Institute", is_active=False,
                                           organization=self.org)
        ClientBilling.objects.create(client=gone, academic_year="2025-2026",
                                     override_amounts=True, net_amount=Decimal("99999"))

    def test_forecast_excludes_discontinued(self):
        from payslip.services.income_analytics import build_forecast
        names = [r["name"] for r in build_forecast(self.org)["rows"]]
        self.assertNotIn("Gone Institute", names)
        self.assertIn("Growth College", names)
        self.assertIn("Fixed School", names)

    def test_forecast_growth_projection(self):
        from decimal import Decimal
        from payslip.services.income_analytics import build_forecast
        fc = build_forecast(self.org)
        self.assertEqual(fc["target_year"], "2026-2027")
        grower = next(r for r in fc["rows"] if r["name"] == "Growth College")
        # 1100 * 1.10 growth = 1210 projected students
        self.assertEqual(grower["projected_count"], 1210)
        # 1210 * 100 * 1.18 = 142780
        self.assertEqual(grower["growth"], Decimal("142780.00"))
        # conservative = latest net (1100*100*1.18 = 129800)
        self.assertEqual(grower["conservative"], Decimal("129800.00"))

    def test_forecast_fixed_fee_uses_same_net(self):
        from decimal import Decimal
        from payslip.services.income_analytics import build_forecast
        fixed = next(r for r in build_forecast(self.org)["rows"] if r["name"] == "Fixed School")
        self.assertEqual(fixed["conservative"], Decimal("350000.00"))
        self.assertEqual(fixed["growth"], Decimal("350000.00"))
        self.assertEqual(fixed["note"], "fixed fee")

    def test_analytics_fy_totals(self):
        from decimal import Decimal
        from payslip.services.income_analytics import build_analytics
        a = build_analytics(self.org)
        fy26 = next(r for r in a["fy_rows"] if r["year"] == "2025-2026")
        # billed = 129800 (grower) + 350000 (fixed) + 99999 (gone)
        self.assertEqual(fy26["billed"], Decimal("579799.00"))
        self.assertEqual(fy26["received"], Decimal("50000.00"))

    def test_analytics_page_renders_for_staff(self):
        resp = self.client.get("/income/analytics/")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Revenue Forecast")
        self.assertContains(resp, "Growth College")
        self.assertNotContains(resp, 'Gone Institute</a></td>\n              <td class="num">₹')

    def test_analytics_forbidden_without_income_right(self):
        plain = make_member("inc_plain2", org=self.org)
        self.client.force_login(plain)
        self.assertEqual(self.client.get("/income/analytics/").status_code, 403)


class ImplementationTrackingTests(TestCase):
    def setUp(self):
        from .models import IncomeClient
        self.staff = make_member("impl_staff", income=True, implementation=True)
        self.org = self.staff.membership.organization
        self.plain = make_member("impl_plain", org=self.org)
        self.income_client = IncomeClient.objects.create(name="Impl College",
                                                         organization=self.org)
        self.client.force_login(self.staff)

    def test_agreement_end_auto_computes_from_start_plus_years(self):
        import datetime
        from .models import ClientOnboarding
        o = ClientOnboarding.objects.create(
            client=self.income_client,
            agreement_start=datetime.date(2025, 6, 1), agreement_years=3,
        )
        self.assertEqual(o.agreement_end, datetime.date(2028, 6, 1))

    def test_explicit_agreement_end_wins(self):
        import datetime
        from .models import ClientOnboarding
        o = ClientOnboarding.objects.create(
            client=self.income_client,
            agreement_start=datetime.date(2025, 6, 1), agreement_years=3,
            agreement_end=datetime.date(2027, 12, 31),
        )
        self.assertEqual(o.agreement_end, datetime.date(2027, 12, 31))

    def test_expiry_properties_across_boundaries(self):
        import datetime
        from django.utils import timezone
        from .models import ClientOnboarding
        today = timezone.localdate()
        o = ClientOnboarding.objects.create(client=self.income_client,
                                            agreement_signed=True)
        # 89 days out: inside the 90-day window -> expiring, not expired
        o.agreement_end = today + datetime.timedelta(days=89)
        self.assertTrue(o.agreement_expiring)
        self.assertFalse(o.agreement_expired)
        # 91 days out: outside the window
        o.agreement_end = today + datetime.timedelta(days=91)
        self.assertFalse(o.agreement_expiring)
        self.assertFalse(o.agreement_expired)
        # past date: expired, not expiring
        o.agreement_end = today - datetime.timedelta(days=1)
        self.assertTrue(o.agreement_expired)
        self.assertFalse(o.agreement_expiring)
        # no end date at all
        o.agreement_end = None
        self.assertIsNone(o.days_to_expiry)
        self.assertFalse(o.agreement_expired)
        self.assertFalse(o.agreement_expiring)

    def test_seed_cms_creates_25_modules_and_is_idempotent(self):
        url = f"/income/clients/{self.income_client.pk}/implementation/"
        resp = self.client.post(url, {"action": "seed_cms"})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(self.income_client.features.count(), 25)
        self.client.post(url, {"action": "seed_cms"})
        self.assertEqual(self.income_client.features.count(), 25)

    def test_progress_excludes_na(self):
        from .models import FeatureStatus, feature_progress
        S = FeatureStatus.Status
        for i, status in enumerate([S.LIVE, S.LIVE, S.IN_PROGRESS, S.NA]):
            FeatureStatus.objects.create(client=self.income_client,
                                         name=f"F{i}", status=status, order=i)
        p = feature_progress(self.income_client.features.all())
        self.assertEqual(p["total"], 4)
        self.assertEqual(p["applicable"], 3)
        self.assertEqual(p["live"], 2)
        self.assertEqual(p["pct"], 66)

    def test_bulk_save_stamps_completed_on_for_live(self):
        from django.utils import timezone
        from .models import FeatureStatus
        f = FeatureStatus.objects.create(client=self.income_client, name="Fees")
        url = f"/income/clients/{self.income_client.pk}/implementation/"
        resp = self.client.post(url, {
            "action": "save_features",
            f"status_{f.pk}": "LIVE",
            f"engineer_{f.pk}": "Kalai",
            f"remarks_{f.pk}": "done",
        })
        self.assertEqual(resp.status_code, 302)
        f.refresh_from_db()
        self.assertEqual(f.status, "LIVE")
        self.assertEqual(f.completed_on, timezone.localdate())
        self.assertEqual(f.engineer, "Kalai")

    def test_access_control_both_routes(self):
        dash = "/income/implementation/"
        detail = f"/income/clients/{self.income_client.pk}/implementation/"
        self.assertEqual(self.client.get(dash).status_code, 200)
        self.assertEqual(self.client.get(detail).status_code, 200)
        self.client.force_login(self.plain)
        self.assertEqual(self.client.get(dash).status_code, 403)
        self.assertEqual(self.client.get(detail).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(dash).status_code, 302)
        self.assertEqual(self.client.get(detail).status_code, 302)

    def test_dashboard_lists_expiring_and_po_pending(self):
        import datetime
        from django.utils import timezone
        from .models import ClientOnboarding
        ClientOnboarding.objects.create(
            client=self.income_client, agreement_signed=True,
            agreement_end=timezone.localdate() + datetime.timedelta(days=30),
        )
        resp = self.client.get("/income/implementation/")
        self.assertContains(resp, "Impl College")
        self.assertContains(resp, "Agreements needing attention")
        self.assertContains(resp, "PO pending")

    def test_income_dashboard_shows_alert_banner(self):
        resp = self.client.get("/income/")
        self.assertContains(resp, "without a PO")


class MembershipProvisioningTests(TestCase):
    def test_signup_creates_org_and_admin_membership(self):
        from payslip.models import Membership
        resp = self.client.post("/accounts/signup/", {
            "username": "neworg_admin", "email": "neworg@x.com",
            "password1": "str0ng-pass-123", "password2": "str0ng-pass-123",
        })
        self.assertEqual(resp.status_code, 302)
        self.assertIn("welcome=1", resp["Location"])
        user = User.objects.get(username="neworg_admin")
        m = Membership.objects.get(user=user)
        self.assertEqual(m.role, Membership.Role.ADMIN)
        self.assertIsNotNone(m.organization)

    def test_membership_for_auto_provisions_bare_user(self):
        from payslip.models import Membership, membership_for
        user = User.objects.create_user("bare_user", "bare@x.com", "pass12345")
        m = membership_for(user)
        self.assertEqual(m.role, Membership.Role.ADMIN)
        self.assertEqual(membership_for(user).pk, m.pk)  # stable on re-call


class ModuleRightsTests(TestCase):
    def setUp(self):
        self.admin = make_member("rights_admin", admin=True)
        self.org = self.admin.membership.organization

    def test_member_without_right_gets_friendly_403(self):
        member = make_member("no_travel", org=self.org, offer_letters=True)
        self.client.force_login(member)
        resp = self.client.get("/travel-expense/")
        self.assertEqual(resp.status_code, 403)
        self.assertContains(resp, "organization admin", status_code=403)

    def test_member_with_right_gets_200(self):
        member = make_member("has_travel", org=self.org, travel_expense=True)
        self.client.force_login(member)
        self.assertEqual(self.client.get("/travel-expense/").status_code, 200)
        # ...but proposals stays blocked
        self.assertEqual(self.client.get("/proposal-quotation/").status_code, 403)

    def test_admin_bypasses_toggles(self):
        from payslip.models import Membership
        Membership.objects.filter(user=self.admin).update(
            can_travel_expense=False, can_income=False)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get("/travel-expense/").status_code, 200)
        self.assertEqual(self.client.get("/income/").status_code, 200)

    def test_anonymous_redirects_to_login(self):
        r = self.client.get("/travel-expense/")
        self.assertEqual(r.status_code, 302)
        self.assertIn("/accounts/login/", r["Location"])


class OrgIsolationTests(TestCase):
    def setUp(self):
        from decimal import Decimal
        from payslip.models import ClientBilling, IncomeClient
        self.user_a = make_member("org_a_user", income=True, implementation=True)
        self.org_a = self.user_a.membership.organization
        self.user_b = make_member("org_b_user", income=True, implementation=True)
        self.org_b = self.user_b.membership.organization
        # Same client name in both orgs - allowed by the composite constraint.
        self.client_a = IncomeClient.objects.create(name="Test College", organization=self.org_a)
        self.client_b = IncomeClient.objects.create(name="Test College", organization=self.org_b)
        ClientBilling.objects.create(client=self.client_a, academic_year="2025-2026",
                                     override_amounts=True, net_amount=Decimal("100000"))
        ClientBilling.objects.create(client=self.client_b, academic_year="2025-2026",
                                     override_amounts=True, net_amount=Decimal("777777"))

    def test_client_list_scoped(self):
        self.client.force_login(self.user_a)
        resp = self.client.get("/income/clients/")
        self.assertContains(resp, "Test College")
        self.assertNotContains(resp, "777,777")

    def test_cross_org_detail_edit_implementation_404(self):
        self.client.force_login(self.user_a)
        for url in (f"/income/clients/{self.client_b.pk}/",
                    f"/income/clients/{self.client_b.pk}/edit/",
                    f"/income/clients/{self.client_b.pk}/implementation/"):
            self.assertEqual(self.client.get(url).status_code, 404, url)

    def test_analytics_scoped(self):
        from decimal import Decimal
        from payslip.services.income_analytics import build_analytics
        a = build_analytics(self.org_a)
        total_billed = sum(r["billed"] for r in a["fy_rows"])
        self.assertEqual(total_billed, Decimal("100000.00"))

    def test_export_scoped(self):
        from openpyxl import load_workbook
        from io import BytesIO
        from payslip.services.income_export import build_income_workbook
        wb = load_workbook(BytesIO(build_income_workbook(self.org_b)))
        values = [str(c.value) for row in wb.active.iter_rows() for c in row if c.value]
        self.assertTrue(any("777777" in v for v in values))
        self.assertFalse(any(v == "100000" for v in values))


class TeamPageTests(TestCase):
    def setUp(self):
        self.admin = make_member("team_admin", admin=True)
        self.org = self.admin.membership.organization
        self.client.force_login(self.admin)

    def test_member_cannot_open_team_page(self):
        member = make_member("team_member", org=self.org)
        self.client.force_login(member)
        self.assertEqual(self.client.get("/team/").status_code, 403)

    def test_admin_adds_member_and_they_can_log_in(self):
        from payslip.models import Membership
        resp = self.client.post("/team/add/", {
            "first_name": "Priya", "username": "priya", "email": "priya@x.com",
            "password": "temp-pass-9x21", "role": "MEMBER",
            "new_can_payroll": "on",
        })
        self.assertEqual(resp.status_code, 302)
        user = User.objects.get(username="priya")
        m = Membership.objects.get(user=user)
        self.assertEqual(m.organization, self.org)
        self.assertTrue(m.can_payroll)
        self.assertFalse(m.can_income)
        # Temp password works for a real login.
        fresh = self.client_class()
        self.assertTrue(fresh.login(username="priya", password="temp-pass-9x21"))

    def test_deactivate_blocks_next_request(self):
        member = make_member("leaver", org=self.org, travel_expense=True)
        member_session = self.client_class()
        member_session.force_login(member)
        self.assertEqual(member_session.get("/travel-expense/").status_code, 200)
        resp = self.client.post(f"/team/{member.pk}/toggle/")
        self.assertEqual(resp.status_code, 302)
        member.refresh_from_db()
        self.assertFalse(member.is_active)
        # Their live session no longer passes module_required.
        self.assertNotEqual(member_session.get("/travel-expense/").status_code, 200)

    def test_bulk_save_updates_rights(self):
        member = make_member("flipme", org=self.org)
        resp = self.client.post("/team/", {
            f"role_{self.admin.pk}": "ADMIN",
            f"can_payroll_{self.admin.pk}": "on",
            f"role_{member.pk}": "MEMBER",
            f"can_income_{member.pk}": "on",
        })
        self.assertEqual(resp.status_code, 302)
        member.membership.refresh_from_db()
        self.assertTrue(member.membership.can_income)
        self.assertFalse(member.membership.can_payroll)

    def test_last_admin_cannot_be_demoted(self):
        from payslip.models import Membership
        resp = self.client.post("/team/", {f"role_{self.admin.pk}": "MEMBER"})
        self.assertEqual(resp.status_code, 302)
        m = Membership.objects.get(user=self.admin)
        self.assertEqual(m.role, Membership.Role.ADMIN)  # unchanged

    def test_last_admin_cannot_be_deactivated(self):
        # Deactivate the founder via a second admin, then verify the last
        # active admin can be neither demoted nor self-deactivated.
        second_admin = make_member("second_admin", admin=True, org=self.org)
        second_admin_session = self.client_class()
        second_admin_session.force_login(second_admin)
        second_admin_session.post(f"/team/{self.admin.pk}/toggle/")
        self.admin.refresh_from_db()
        self.assertFalse(self.admin.is_active)
        # Demote via bulk save: blocked by the active-admin count.
        second_admin_session.post("/team/", {f"role_{second_admin.pk}": "MEMBER"})
        second_admin.membership.refresh_from_db()
        self.assertEqual(second_admin.membership.role, "ADMIN")
        # Self-deactivation: blocked by the self-guard.
        second_admin_session.post(f"/team/{second_admin.pk}/toggle/")
        second_admin.refresh_from_db()
        self.assertTrue(second_admin.is_active)


class ProposalHistoryTests(TestCase):
    def setUp(self):
        self.user_a = make_member("prop_a", admin=True)
        self.org_a = self.user_a.membership.organization
        self.user_b = make_member("prop_b", admin=True)
        self.client.force_login(self.user_a)

    def _generate(self, client_name="ABC College of Arts and Science", **overrides):
        payload = {**PROPOSAL_PAYLOAD, "client_name": client_name, **overrides}
        return self.client.post(reverse("proposal_quotation"), payload)

    def test_revision_chain_per_client(self):
        from payslip.models import ProposalRecord
        self._generate()
        self._generate()
        self._generate(client_name="XYZ School")
        revisions = list(
            ProposalRecord.objects.filter(client_name="ABC College of Arts and Science")
            .order_by("revision").values_list("revision", flat=True)
        )
        self.assertEqual(revisions, [1, 2])
        xyz = ProposalRecord.objects.get(client_name="XYZ School")
        self.assertEqual(xyz.revision, 1)

    def test_history_list_and_cross_org_isolation(self):
        from payslip.models import ProposalRecord
        self._generate()
        record = ProposalRecord.objects.get()

        resp = self.client.get(reverse("proposal_history"))
        self.assertContains(resp, "ABC College of Arts and Science")
        self.assertContains(resp, "Rev 1")

        # Org B sees an empty history and 404s on org A's record.
        self.client.force_login(self.user_b)
        resp = self.client.get(reverse("proposal_history"))
        self.assertNotContains(resp, "ABC College of Arts and Science")
        self.assertEqual(
            self.client.get(reverse("proposal_record", args=[record.pk])).status_code, 404
        )

    def test_revise_prefills_form(self):
        from payslip.models import ProposalRecord
        self._generate()
        record = ProposalRecord.objects.get()
        resp = self.client.get(f"{reverse('proposal_quotation')}?from={record.pk}")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Revising a proposal from history")
        self.assertContains(resp, 'value="ABC College of Arts and Science"')
        self.assertContains(resp, 'value="13/02/2026"')

    def test_view_action_serves_stored_html(self):
        from payslip.models import ProposalRecord
        self._generate()
        record = ProposalRecord.objects.get()
        resp = self.client.post(reverse("proposal_record", args=[record.pk]),
                                {"action": "view"})
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/preview/", resp["Location"])
        preview = self.client.get(resp["Location"])
        self.assertEqual(preview.status_code, 200)
        self.assertIn(b"ABC College of Arts and Science", preview.content)

    def test_download_action_returns_file(self):
        from payslip.models import ProposalRecord
        self._generate()
        record = ProposalRecord.objects.get()
        resp = self.client.post(reverse("proposal_record", args=[record.pk]),
                                {"action": "download"})
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/download/", resp["Location"])
        download = self.client.get(resp["Location"])
        self.assertEqual(download.status_code, 200)
        self.assertIn("attachment", download["Content-Disposition"])

    def test_delete_action_removes_record(self):
        from payslip.models import ProposalRecord
        self._generate()
        record = ProposalRecord.objects.get()
        resp = self.client.post(reverse("proposal_record", args=[record.pk]),
                                {"action": "delete"})
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(ProposalRecord.objects.count(), 0)

    def test_record_page_shows_revision_chain(self):
        from payslip.models import ProposalRecord
        self._generate()
        self._generate()
        rev2 = ProposalRecord.objects.get(revision=2)
        resp = self.client.get(reverse("proposal_record", args=[rev2.pk]))
        self.assertContains(resp, "Rev 1")
        self.assertContains(resp, "Rev 2")
        self.assertContains(resp, "Revision history")

    def test_member_without_proposals_right_gets_403(self):
        member = make_member("no_props", org=self.org_a)  # all rights False
        self.client.force_login(member)
        self.assertEqual(self.client.get(reverse("proposal_history")).status_code, 403)


# ONE_TIME + AMC commercial-model payload (per-student fields intentionally
# left out - the mode must not require them).
ONE_TIME_PAYLOAD = {
    "to_address": "The Principal",
    "client_name": "AMC Test College",
    "client_address": "Coimbatore, Tamil Nadu",
    "proposal_date": "13/02/2026",
    "prepared_by": "Aveon Infotech Private Limited",
    "selection_mode": "BUNDLE",
    "bundle": "CMS_FULL",
    "pricing_model": "ONE_TIME",
    "one_time_price": "1000000",
    "amc_percent": "18",
    "amc_amount": "180000",
    "one_time_implementation_fee": "350000",
    "gst_percent": "18",
    "authorized_signatory_name": "Parvathi G",
    "authorized_signatory_designation": "Chief Executive Officer",
}

AMC_TERM = "The AMC shall remain fixed for the first three years."
LICENSE_TERM = "The license fee shall remain fixed for the first three years."


class ProposalPricingModelTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user("amc_tester", "amc@example.com", "pass12345")
        self.client.force_login(self.user)

    def test_one_time_generates_with_amc_lines_and_terms(self):
        from decimal import Decimal
        from payslip.models import ProposalRecord
        response = self.client.post(reverse("proposal_quotation"), ONE_TIME_PAYLOAD)
        self.assertEqual(response.status_code, 200)
        record = ProposalRecord.objects.get()
        # Year-1 total: (10,00,000 + 3,50,000) x 1.18 (fee not waived)
        self.assertEqual(record.total_amount, Decimal("1593000.00"))
        self.assertIn("One-time + AMC", record.selection_label)
        self.assertIn("One-Time License Fee", record.html)
        self.assertIn("Annual Maintenance Contract", record.html)
        self.assertIn("from Year 2 onwards", record.html)
        self.assertIn(AMC_TERM, record.html)
        self.assertNotIn(LICENSE_TERM, record.html)

    def test_per_student_keeps_license_term(self):
        from payslip.models import ProposalRecord
        self.client.post(reverse("proposal_quotation"), PROPOSAL_PAYLOAD)
        record = ProposalRecord.objects.get()
        self.assertIn(LICENSE_TERM, record.html)
        self.assertNotIn(AMC_TERM, record.html)
        self.assertNotIn("Annual Maintenance Contract", record.html)
        self.assertNotIn("One-time + AMC", record.selection_label)

    def test_amc_amount_computed_from_percent_when_blank(self):
        from payslip.models import ProposalRecord
        payload = {**ONE_TIME_PAYLOAD, "amc_amount": ""}
        response = self.client.post(reverse("proposal_quotation"), payload)
        self.assertEqual(response.status_code, 200)
        record = ProposalRecord.objects.get()
        self.assertEqual(record.form_data["amc_amount"], "180000.00")
        self.assertIn("1,80,000", record.html)

    def test_one_time_requires_price_and_amc(self):
        from payslip.models import ProposalRecord
        payload = {**ONE_TIME_PAYLOAD, "one_time_price": "",
                   "amc_percent": "", "amc_amount": ""}
        response = self.client.post(reverse("proposal_quotation"), payload)
        self.assertEqual(response.status_code, 200)
        errors = response.context["form"].errors
        self.assertIn("one_time_price", errors)
        self.assertIn("amc_amount", errors)
        self.assertEqual(ProposalRecord.objects.count(), 0)

    def test_per_student_mode_still_requires_its_fields(self):
        from payslip.models import ProposalRecord
        payload = {**PROPOSAL_PAYLOAD, "price_per_unit": "",
                   "minimum_student_commitment": ""}
        response = self.client.post(reverse("proposal_quotation"), payload)
        self.assertEqual(response.status_code, 200)
        errors = response.context["form"].errors
        self.assertIn("price_per_unit", errors)
        self.assertIn("minimum_student_commitment", errors)
        self.assertEqual(ProposalRecord.objects.count(), 0)

    def test_revise_prefills_one_time_mode(self):
        from payslip.models import ProposalRecord
        self.client.post(reverse("proposal_quotation"), ONE_TIME_PAYLOAD)
        record = ProposalRecord.objects.get()
        resp = self.client.get(f"{reverse('proposal_quotation')}?from={record.pk}")
        self.assertEqual(resp.status_code, 200)
        self.assertRegex(resp.content.decode(), r'value="ONE_TIME"\s+checked')
        self.assertContains(resp, 'value="1000000"')


INTERNSHIP_OFFER_PAYLOAD = {
    "offer_type": "internship",
    "name": "Kavya R",
    "roll_number": "21CS042",
    "course": "B.E. CSE",
    "college_name": "ABC Engineering College",
    "college_address": "Coimbatore",
    "internship_role": "Full Stack Developer",
    "start_date": "2026-08-01",
    "duration_months": "3",
    "intern_signatory": "Ranjith Kumar",
    "intern_signatory_designation": "General Manager",
}

APPOINTMENT_PAYLOAD = {
    "offer_type": "appointment",
    "serial_no": "AV/2026/014",
    "employee_name": "Suresh Kumar",
    "designation": "Software Engineer",
    "join_date": "2026-08-10",
    "company_name": "Aveon Infotech Private Limited",
    "signatory": "Parvathi G",
    "signatory_designation": "CEO",
}

EXPERIENCE_EMPLOYEE_PAYLOAD = {
    "certificate_type": "employee",
    "gender": "male",
    "title": "Mr.",
    "employee_name_exp": "Suresh Kumar",
    "employee_no": "AV104",
    "company_name_exp": "Aveon Infotech Private Limited",
    "join_date_exp": "2023-06-01",
    "leaving_date": "2026-06-30",
    "designation_exp": "Senior Software Engineer",
    "signatory_exp": "Parvathi G",
    "signatory_designation_exp": "CEO",
}


class PeopleRegistryTests(TestCase):
    def setUp(self):
        self.user = make_member("people_admin", admin=True)
        self.org = self.user.membership.organization
        self.client.force_login(self.user)

    def test_internship_offer_auto_creates_intern(self):
        from payslip.models import Person, PersonDocument
        resp = self.client.post(reverse("offer_letter"), INTERNSHIP_OFFER_PAYLOAD)
        self.assertEqual(resp.status_code, 200)
        person = Person.objects.get()
        self.assertEqual(person.kind, Person.Kind.INTERN)
        self.assertEqual(person.name, "Kavya R")
        self.assertEqual(person.college_name, "ABC Engineering College")
        self.assertEqual(person.organization, self.org)
        doc = PersonDocument.objects.get()
        self.assertEqual(doc.doc_type, PersonDocument.DocType.INTERNSHIP_OFFER)
        self.assertTrue(bytes(doc.pdf).startswith(b"%PDF-"))
        self.assertTrue(bytes(doc.pdf_plain).startswith(b"%PDF-"))
        # Regenerating for the same name reuses the person.
        self.client.post(reverse("offer_letter"), INTERNSHIP_OFFER_PAYLOAD)
        self.assertEqual(Person.objects.count(), 1)
        self.assertEqual(person.documents.count(), 2)

    def test_appointment_creates_candidate_and_experience_updates(self):
        import datetime
        from payslip.models import Person
        self.client.post(reverse("offer_letter"), APPOINTMENT_PAYLOAD)
        person = Person.objects.get()
        self.assertEqual(person.kind, Person.Kind.CANDIDATE)
        self.assertEqual(person.designation, "Software Engineer")
        self.assertEqual(person.join_date, datetime.date(2026, 8, 10))
        # An employee experience letter for the same name updates the record.
        self.client.post(reverse("experience_certificate"), EXPERIENCE_EMPLOYEE_PAYLOAD)
        person.refresh_from_db()
        self.assertEqual(Person.objects.count(), 1)
        self.assertEqual(person.leaving_date, datetime.date(2026, 6, 30))
        self.assertEqual(person.designation, "Senior Software Engineer")
        self.assertEqual(person.documents.count(), 2)

    def test_person_prefill_offer_letter(self):
        from payslip.models import Person
        self.client.post(reverse("offer_letter"), INTERNSHIP_OFFER_PAYLOAD)
        person = Person.objects.get()
        resp = self.client.get(f"{reverse('offer_letter')}?person={person.pk}&type=internship")
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Prefilled from Kavya R")
        self.assertContains(resp, 'value="Kavya R"')
        self.assertContains(resp, 'value="21CS042"')
        # A prefilled page must tell its JS to skip the localStorage draft
        # restore - otherwise a stale empty draft silently wipes the name
        # back out (the bug: "candidate name not loading").
        self.assertContains(resp, "const HAS_PREFILL = true;")

    def test_blank_offer_letter_page_does_not_skip_draft_restore(self):
        resp = self.client.get(reverse("offer_letter"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "const HAS_PREFILL = false;")

    def test_cross_org_prefill_and_pages_404(self):
        from payslip.models import Person
        self.client.post(reverse("offer_letter"), INTERNSHIP_OFFER_PAYLOAD)
        person = Person.objects.get()
        other = make_member("other_org_admin", admin=True)
        self.client.force_login(other)
        self.assertEqual(
            self.client.get(f"{reverse('offer_letter')}?person={person.pk}&type=internship").status_code,
            404,
        )
        self.assertEqual(
            self.client.get(reverse("person_detail", args=[person.pk])).status_code, 404
        )
        resp = self.client.get(reverse("people_list"))
        self.assertNotContains(resp, "Kavya R")

    def test_people_rights_gating(self):
        member = make_member("no_people", org=self.org, offer_letters=True)
        self.client.force_login(member)
        self.assertEqual(self.client.get(reverse("people_list")).status_code, 403)
        allowed = make_member("has_people", org=self.org, people=True)
        self.client.force_login(allowed)
        self.assertEqual(self.client.get(reverse("people_list")).status_code, 200)

    def test_download_delete_doc_and_person(self):
        from payslip.models import Person, PersonDocument
        self.client.post(reverse("offer_letter"), INTERNSHIP_OFFER_PAYLOAD)
        person = Person.objects.get()
        doc = person.documents.first()
        url = reverse("person_detail", args=[person.pk])

        resp = self.client.post(url, {"action": "download_doc", "doc_id": doc.pk})
        self.assertEqual(resp.status_code, 302)
        self.assertIn("/download/", resp["Location"])
        dl = self.client.get(resp["Location"])
        self.assertEqual(dl.status_code, 200)
        self.assertTrue(dl.content.startswith(b"%PDF-"))

        self.client.post(url, {"action": "delete_doc", "doc_id": doc.pk})
        self.assertEqual(PersonDocument.objects.count(), 0)

        self.client.post(url, {"action": "delete_person"})
        self.assertEqual(Person.objects.count(), 0)

    def test_manual_person_create_and_edit(self):
        from payslip.models import Person
        resp = self.client.post(reverse("person_create"), {
            "kind": "CANDIDATE", "name": "Manual Candidate",
            "designation": "Analyst",
        })
        self.assertEqual(resp.status_code, 302)
        person = Person.objects.get()
        self.assertEqual(person.organization, self.org)
        resp = self.client.post(reverse("person_detail", args=[person.pk]), {
            "action": "save_person", "kind": "CANDIDATE",
            "name": "Manual Candidate", "designation": "Senior Analyst",
        })
        self.assertEqual(resp.status_code, 302)
        person.refresh_from_db()
        self.assertEqual(person.designation, "Senior Analyst")

    def test_people_list_splits_candidates_and_interns(self):
        self.client.post(reverse("offer_letter"), APPOINTMENT_PAYLOAD)
        self.client.post(reverse("offer_letter"), INTERNSHIP_OFFER_PAYLOAD)
        resp = self.client.get(reverse("people_list"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Candidates / Employees")
        self.assertContains(resp, "Interns")
        self.assertContains(resp, "Suresh Kumar")   # candidate
        self.assertContains(resp, "Kavya R")         # intern

    def test_quick_generate_links_present(self):
        resp = self.client.get(reverse("people_list"))
        self.assertContains(resp, 'href="/offer-letter/"')
        self.assertContains(resp, 'href="/experience-certificate/"')

    def test_landing_hides_letter_menu_items_for_authenticated_user(self):
        member = make_member("landing_check", org=self.org,
                             offer_letters=True, experience_certificates=True,
                             people=True)
        self.client.force_login(member)
        resp = self.client.get(reverse("landing"))
        self.assertEqual(resp.status_code, 200)
        self.assertNotContains(resp, 'href="/offer-letter/"')
        self.assertNotContains(resp, 'href="/experience-certificate/"')
        self.assertContains(resp, "People")

    def test_landing_shows_letter_cards_for_anonymous(self):
        self.client.logout()
        resp = self.client.get(reverse("landing"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Offer Letter Generator")
        self.assertContains(resp, "Experience Certificate Generator")
        self.assertContains(resp, 'href="/offer-letter/"')


class PayrollCalcTests(TestCase):
    """Verified against the real source Salary Excel (Vijayalakshmi and
    Raja.S rows, cross-checked cell-by-cell via openpyxl data_only=True)."""

    def _settings(self):
        from payslip.models import PayrollSettings
        return PayrollSettings()  # unsaved instance - defaults match the sheet

    def test_full_attendance_high_salary_not_esi_eligible(self):
        from decimal import Decimal
        from payslip.services.payroll_calc import compute_entry
        result = compute_entry(
            monthly_package=Decimal("37000"),
            total_working_days=30,
            emp_leave_days=Decimal("3"),
            lop_days=Decimal("0"),
            is_esi_eligible=False,
            is_pf_applicable=True,
            settings=self._settings(),
        )
        self.assertEqual(result.present_days, Decimal("27"))
        self.assertEqual(result.pay_days, Decimal("30"))
        self.assertEqual(result.basic, Decimal("18500"))
        self.assertEqual(result.da, Decimal("8325.00"))
        self.assertEqual(result.hra, Decimal("4625.00"))
        self.assertEqual(result.transport_allowance, Decimal("3700.00"))
        self.assertEqual(result.food_allowance, Decimal("1850.00"))
        self.assertEqual(result.gross_salary, Decimal("37000.00"))
        self.assertEqual(result.esi_employee, Decimal("0"))
        self.assertEqual(result.esi_employer, Decimal("0"))
        # PF wage base (18500+8325)*0.6 = 16095, capped at 15000.
        self.assertEqual(result.pf_employee, Decimal("1800.00"))
        self.assertEqual(result.pf_employer, Decimal("1800.00"))
        self.assertEqual(result.total_deductions, Decimal("1800.00"))
        self.assertEqual(result.net_payable, Decimal("35200.00"))

    def test_esi_eligible_with_arrear_and_uncapped_pf(self):
        from decimal import Decimal
        from payslip.services.payroll_calc import compute_entry
        result = compute_entry(
            monthly_package=Decimal("20000"),
            total_working_days=30,
            emp_leave_days=Decimal("0"),
            lop_days=Decimal("0"),
            salary_arrear_allowance=Decimal("5000"),
            is_esi_eligible=True,
            is_pf_applicable=True,
            settings=self._settings(),
        )
        self.assertEqual(result.basic, Decimal("10000"))
        self.assertEqual(result.da, Decimal("4500.00"))
        self.assertEqual(result.hra, Decimal("2500.00"))
        self.assertEqual(result.transport_allowance, Decimal("2000.00"))
        self.assertEqual(result.food_allowance, Decimal("1000.00"))
        # Gross includes the one-off arrear; ESI base does not.
        self.assertEqual(result.gross_salary, Decimal("25000.00"))
        self.assertEqual(result.esi_employee, Decimal("150"))
        self.assertEqual(result.esi_employer, Decimal("650"))
        # PF wage base (10000+4500)*0.6 = 8700, uncapped.
        self.assertEqual(result.pf_employee, Decimal("1044.00"))
        self.assertEqual(result.pf_employer, Decimal("1044.00"))
        self.assertEqual(result.total_deductions, Decimal("1194.00"))
        self.assertEqual(result.net_payable, Decimal("23806.00"))

    def test_esi_eligibility_is_a_flag_not_derived_from_gross(self):
        """A large one-off arrear can push gross above the ESI wage ceiling
        without losing coverage - eligibility must be passed in, not
        recomputed from that month's gross (confirmed against Raja.S)."""
        from decimal import Decimal
        from payslip.services.payroll_calc import compute_entry
        settings = self._settings()
        eligible = compute_entry(
            monthly_package=Decimal("20000"), total_working_days=30,
            salary_arrear_allowance=Decimal("5000"),  # gross 25000 > 21000 ceiling
            is_esi_eligible=True, settings=settings,
        )
        not_eligible = compute_entry(
            monthly_package=Decimal("20000"), total_working_days=30,
            salary_arrear_allowance=Decimal("5000"),
            is_esi_eligible=False, settings=settings,
        )
        self.assertGreater(eligible.esi_employee, Decimal("0"))
        self.assertEqual(not_eligible.esi_employee, Decimal("0"))

    def test_pf_applicability_is_a_flag_not_universal(self):
        """Confirmed against the real sheet: only about half the roster is
        ever PF-enrolled - the PF Employee cell is a literal 0 (never a
        formula) for everyone else. PF must be gated the same way ESI is."""
        from decimal import Decimal
        from payslip.services.payroll_calc import compute_entry
        settings = self._settings()
        applicable = compute_entry(
            monthly_package=Decimal("18000"), total_working_days=30,
            is_pf_applicable=True, settings=settings,
        )
        not_applicable = compute_entry(
            monthly_package=Decimal("18000"), total_working_days=30,
            is_pf_applicable=False, settings=settings,
        )
        self.assertGreater(applicable.pf_employee, Decimal("0"))
        self.assertEqual(applicable.pf_employer, applicable.pf_employee)
        self.assertEqual(not_applicable.pf_employee, Decimal("0"))
        self.assertEqual(not_applicable.pf_employer, Decimal("0"))
        # PF being off doesn't touch gross - only deductions/net.
        self.assertEqual(applicable.gross_salary, not_applicable.gross_salary)

    def test_partial_attendance_prorates_basic(self):
        from decimal import Decimal
        from payslip.services.payroll_calc import compute_entry
        result = compute_entry(
            monthly_package=Decimal("18500"), total_working_days=30,
            lop_days=Decimal("2"),  # pay_days = 28
            is_esi_eligible=False, settings=self._settings(),
        )
        self.assertEqual(result.pay_days, Decimal("28"))
        # ROUND(((18500*0.5)/30)*28, 0) = ROUND(8633.33..., 0) = 8633
        self.assertEqual(result.basic, Decimal("8633"))

    def test_gross_equals_package_at_full_attendance_no_extras(self):
        """DA%+HRA%+Transport%+Food% always sum to 100% of Basic, so Gross
        equals the monthly package exactly whenever attendance is full and
        there's no internet/arrear allowance - a structural invariant."""
        from decimal import Decimal
        from payslip.services.payroll_calc import compute_entry
        for package in (Decimal("15000"), Decimal("22000"), Decimal("45000")):
            result = compute_entry(
                monthly_package=package, total_working_days=31,
                is_esi_eligible=False, settings=self._settings(),
            )
            self.assertEqual(result.gross_salary, package.quantize(Decimal("0.01")))


import datetime
from decimal import Decimal


class PayrollModuleTests(TestCase):
    def setUp(self):
        self.admin = make_member("payroll_admin", admin=True, payroll=True)
        self.org = self.admin.membership.organization
        self.client.force_login(self.admin)

    def _make_employee(self, **overrides):
        from payslip.models import Employee
        defaults = dict(organization=self.org, employee_code="EMP-0001",
                        name="Test Employee", current_monthly_package=Decimal("30000"))
        defaults.update(overrides)
        return Employee.objects.create(**defaults)

    # --- model constraints -------------------------------------------------
    def test_employee_code_unique_per_org(self):
        from django.db import IntegrityError
        self._make_employee()
        with self.assertRaises(IntegrityError):
            self._make_employee()

    def test_payroll_run_period_clamped_and_unique_per_org(self):
        from payslip.models import PayrollRun
        run = PayrollRun.objects.create(organization=self.org,
                                        period=datetime.date(2026, 7, 15))
        self.assertEqual(run.period, datetime.date(2026, 7, 1))
        from django.db import IntegrityError
        with self.assertRaises(IntegrityError):
            PayrollRun.objects.create(organization=self.org, period=datetime.date(2026, 7, 1))

    def test_payslip_entry_unique_per_run_and_employee(self):
        from payslip.models import PayrollRun, PayslipEntry
        employee = self._make_employee()
        run = PayrollRun.objects.create(organization=self.org, period=datetime.date(2026, 7, 1))
        PayslipEntry.objects.create(organization=self.org, run=run, employee=employee,
                                    monthly_package=Decimal("30000"), total_working_days=31)
        from django.db import IntegrityError
        with self.assertRaises(IntegrityError):
            PayslipEntry.objects.create(organization=self.org, run=run, employee=employee,
                                        monthly_package=Decimal("30000"), total_working_days=31)

    def test_employee_with_payslip_history_cannot_be_deleted(self):
        from django.db.models import ProtectedError
        from payslip.models import PayrollRun, PayslipEntry
        employee = self._make_employee()
        run = PayrollRun.objects.create(organization=self.org, period=datetime.date(2026, 7, 1))
        PayslipEntry.objects.create(organization=self.org, run=run, employee=employee,
                                    monthly_package=Decimal("30000"), total_working_days=31)
        with self.assertRaises(ProtectedError):
            employee.delete()

    # --- access control ------------------------------------------------------
    def test_member_without_payroll_right_gets_403(self):
        member = make_member("payroll_no_right", org=self.org, travel_expense=True)
        self.client.force_login(member)
        self.assertEqual(self.client.get(reverse("payroll_run_list")).status_code, 403)
        self.assertEqual(self.client.get(reverse("employee_list")).status_code, 403)

    def test_member_with_right_gets_200(self):
        member = make_member("payroll_has_right", org=self.org, payroll=True)
        self.client.force_login(member)
        self.assertEqual(self.client.get(reverse("payroll_run_list")).status_code, 200)

    def test_settings_and_finalize_require_org_admin(self):
        member = make_member("payroll_member_only", org=self.org, payroll=True)
        self.client.force_login(member)
        self.assertEqual(self.client.get(reverse("payroll_settings")).status_code, 403)

    def test_cross_org_employee_and_run_404(self):
        employee = self._make_employee()
        other_admin = make_member("payroll_other_org", admin=True, payroll=True)
        self.client.force_login(other_admin)
        self.assertEqual(
            self.client.get(reverse("employee_detail", args=[employee.pk])).status_code, 404
        )

    def test_anonymous_redirects_to_login(self):
        self.client.logout()
        r = self.client.get(reverse("payroll_run_list"))
        self.assertEqual(r.status_code, 302)
        self.assertIn("/accounts/login/", r["Location"])

    # --- run lifecycle -------------------------------------------------------
    def test_creating_run_twice_redirects_into_same_run(self):
        from payslip.models import PayrollRun
        self._make_employee()
        resp1 = self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        resp2 = self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        self.assertEqual(resp1["Location"], resp2["Location"])
        self.assertEqual(PayrollRun.objects.filter(organization=self.org).count(), 1)

    def test_run_create_bulk_populates_computed_defaults(self):
        from payslip.models import PayslipEntry
        self._make_employee(current_monthly_package=Decimal("37000"))
        self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        entry = PayslipEntry.objects.get()
        self.assertEqual(entry.basic, Decimal("18500"))
        self.assertEqual(entry.gross_salary, Decimal("37000.00"))

    def test_finalize_blocks_on_negative_net_without_confirmation(self):
        from payslip.models import PayrollRun
        self._make_employee(current_monthly_package=Decimal("10000"))
        self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        run = PayrollRun.objects.get()
        entry = run.entries.get()
        # A huge TDS forces net_payable negative.
        self.client.post(reverse("payroll_run_detail", args=[run.pk]), {
            "action": "save_grid",
            f"twd_{entry.pk}": "31", f"cl_{entry.pk}": "0", f"empleave_{entry.pk}": "0",
            f"lop_{entry.pk}": "0", f"internet_{entry.pk}": "0", f"arrear_{entry.pk}": "0",
            f"advance_{entry.pk}": "0", f"tds_{entry.pk}": "50000", f"remarks_{entry.pk}": "",
        })
        self.client.post(reverse("payroll_run_finalize", args=[run.pk]))
        run.refresh_from_db()
        self.assertEqual(run.status, PayrollRun.Status.DRAFT)

        self.client.post(reverse("payroll_run_finalize", args=[run.pk]), {"confirm_negative": "1"})
        run.refresh_from_db()
        self.assertEqual(run.status, PayrollRun.Status.FINALIZED)
        entry.refresh_from_db()
        self.assertTrue(bytes(entry.pdf).startswith(b"%PDF-"))

    def test_reopen_clears_pdf_and_reverts_to_draft(self):
        from payslip.models import PayrollRun
        self._make_employee()
        self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        run = PayrollRun.objects.get()
        self.client.post(reverse("payroll_run_finalize", args=[run.pk]))
        run.refresh_from_db()
        self.assertEqual(run.status, PayrollRun.Status.FINALIZED)

        self.client.post(reverse("payroll_run_reopen", args=[run.pk]))
        run.refresh_from_db()
        entry = run.entries.get()
        self.assertEqual(run.status, PayrollRun.Status.DRAFT)
        self.assertIsNone(entry.pdf)

    def test_finalized_run_rejects_grid_edits(self):
        from payslip.models import PayrollRun
        self._make_employee(current_monthly_package=Decimal("37000"))
        self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        run = PayrollRun.objects.get()
        entry = run.entries.get()
        self.client.post(reverse("payroll_run_finalize", args=[run.pk]))
        self.client.post(reverse("payroll_run_detail", args=[run.pk]), {
            "action": "save_grid", f"twd_{entry.pk}": "1",
        })
        entry.refresh_from_db()
        self.assertNotEqual(entry.total_working_days, 1)  # ignored - run is finalized

    # --- PDF / export ---------------------------------------------------------
    def test_entry_pdf_download_serves_valid_pdf(self):
        self._make_employee(current_monthly_package=Decimal("37000"))
        self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        from payslip.models import PayrollRun
        run = PayrollRun.objects.get()
        self.client.post(reverse("payroll_run_finalize", args=[run.pk]))
        entry = run.entries.get()
        resp = self.client.get(reverse("payroll_entry_pdf", args=[entry.pk, "download"]))
        self.assertEqual(resp.status_code, 302)
        download = self.client.get(resp["Location"])
        self.assertEqual(download.status_code, 200)
        self.assertTrue(download.content.startswith(b"%PDF-"))

    def test_register_export_returns_valid_workbook(self):
        from openpyxl import load_workbook
        from io import BytesIO
        self._make_employee()
        self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        from payslip.models import PayrollRun
        run = PayrollRun.objects.get()
        resp = self.client.get(reverse("payroll_register_export", args=[run.pk]))
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(BytesIO(resp.content))
        ws = wb.active
        # Row 1 is the merged title banner, row 3 the headers, data from row 4.
        self.assertEqual(ws["A1"].value, "Salary Statement For The Month Of August 2026")
        self.assertEqual(ws["A3"].value, "S.No")
        self.assertEqual(ws["A4"].value, 1)                  # serial number
        self.assertIn("Test Employee", [c.value for c in ws["B"]])

    def test_reopen_survives_missing_finalized_at(self):
        """Imported runs are FINALIZED with no finalized_at/by - reopening
        one must not blow up formatting those None values."""
        from payslip.models import PayrollRun
        self._make_employee()
        self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        run = PayrollRun.objects.get()
        PayrollRun.objects.filter(pk=run.pk).update(
            status=PayrollRun.Status.FINALIZED, finalized_at=None, finalized_by=None)

        resp = self.client.post(reverse("payroll_run_reopen", args=[run.pk]))
        self.assertEqual(resp.status_code, 302)
        run.refresh_from_db()
        self.assertEqual(run.status, PayrollRun.Status.DRAFT)
        self.assertIn("was finalized on import", run.notes)

    def test_generate_payslips_fills_missing_pdfs_on_finalized_run(self):
        from payslip.models import PayrollRun, PayslipEntry
        self._make_employee()
        self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        run = PayrollRun.objects.get()
        PayrollRun.objects.filter(pk=run.pk).update(status=PayrollRun.Status.FINALIZED)
        self.assertIsNone(PayslipEntry.objects.get().pdf)

        resp = self.client.post(reverse("payroll_generate_payslips", args=[run.pk]))
        self.assertEqual(resp.status_code, 302)
        entry = PayslipEntry.objects.get()
        self.assertTrue(bytes(entry.pdf).startswith(b"%PDF"))
        self.assertIsNotNone(entry.pdf_generated_at)

    def test_comparison_shows_previous_month_movement(self):
        from payslip.models import PayrollRun, PayslipEntry
        self._make_employee()
        self.client.post(reverse("payroll_run_create"), {"period": "2026-07"})
        self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        july, august = PayrollRun.objects.order_by("period")
        # Drop July's net so August shows a rise.
        PayslipEntry.objects.filter(run=july).update(net_payable=Decimal("1000"))

        resp = self.client.get(reverse("payroll_run_detail", args=[august.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Compared with July 2026")
        self.assertEqual(resp.context["comparison"]["prev_run"], july)
        self.assertEqual(len(resp.context["comparison"]["rows"]), 1)

    def test_run_list_is_newest_first(self):
        """annotate() silently drops Meta.ordering, so the run list has to
        restate it - otherwise the 'latest run' KPI shows the oldest month."""
        self._make_employee()
        for period in ("2026-06", "2026-08", "2026-07"):
            self.client.post(reverse("payroll_run_create"), {"period": period})
        resp = self.client.get(reverse("payroll_run_list"))
        periods = [r.period for r in resp.context["runs"]]
        self.assertEqual(periods, sorted(periods, reverse=True))
        self.assertEqual(resp.context["latest_run"].period, datetime.date(2026, 8, 1))

    def test_comparison_absent_for_earliest_run(self):
        from payslip.models import PayrollRun
        self._make_employee()
        self.client.post(reverse("payroll_run_create"), {"period": "2026-08"})
        run = PayrollRun.objects.get()
        resp = self.client.get(reverse("payroll_run_detail", args=[run.pk]))
        self.assertIsNone(resp.context["comparison"])


class ImportPayrollCommandTests(TestCase):
    """Dry-run + commit against a small synthetic in-memory workbook -
    proves the parser, dedup, and cross-validation logic without needing
    the real (external, not-in-repo) Salary_25_26.xlsx."""

    def _build_workbook(self, tmp_path):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        headers = ["Month", "S.No.", "Name", "Designation", "Total Working Days",
                  "CL Credit", "Emp Leave Days", "LOP Days", "Present Days",
                  "Pay Days", "New Salary", "Basic", "DA", "HRA", "Transport",
                  "Food", "Internet", "Arrear", "Gross", "ESI Emp", "ESI Er",
                  "PF Emp", "PF Er", "Advance", "TDS", "Total Ded", "Net"]
        ws.append(headers)
        period = datetime.date(2026, 4, 1)
        # A clean row matching the engine exactly (full attendance, PF+ESI off).
        ws.append([period, 1, "Import Test One", "Dev", 30, 1, 0, 0, 30, 30,
                  20000, 10000, 4500, 2500, 2000, 1000, None, None, 20000,
                  0, 0, 0, 0, 0, 0, 0, 20000])
        # A blank-name row - must be skipped.
        ws.append([period, 2, None, "Dev", 30, 0, 0, 0, 30, 30,
                  0, 0, 0, 0, 0, 0, None, None, 0, 0, 0, 0, 0, 0, 0, 0, 0])
        # A duplicate (same period+name) - first row must win.
        ws.append([period, 1, "Import Test One", "Dev", 30, 1, 0, 5, 30, 25,
                  20000, 8333, 3750, 2083, 1667, 833, None, None, 16666,
                  0, 0, 0, 0, 0, 0, 0, 16666])

        ws4 = wb.create_sheet("Sheet4")
        ws4["B1"] = "S.No"; ws4["C1"] = "Name"; ws4["D1"] = "Current Salary"
        ws4["H1"] = "DOJ"; ws4["K1"] = "Remarks"
        ws4["B2"] = 1; ws4["C2"] = "Import Test One"
        ws4["H2"] = datetime.date(2020, 1, 1)

        path = tmp_path / "test_salary.xlsx"
        wb.save(str(path))
        return str(path)

    def test_dry_run_parses_dedups_and_reports(self):
        import tempfile
        from pathlib import Path
        from io import StringIO
        from django.core.management import call_command

        make_member("import_dry_admin", admin=True)
        with tempfile.TemporaryDirectory() as tmp:
            path = self._build_workbook(Path(tmp))
            out = StringIO()
            call_command("import_payroll", path, username="import_dry_admin", stdout=out)
            output = out.getvalue()
            self.assertIn("Parsed 1 unique", output)  # dedup collapsed 2 rows to 1
            self.assertIn("Dry run only", output)

        from payslip.models import Employee
        self.assertEqual(Employee.objects.count(), 0)  # dry run - nothing written

    def test_commit_creates_employee_run_and_entry_with_dedup_winner(self):
        import tempfile
        from pathlib import Path
        from django.core.management import call_command

        make_member("import_commit_admin", admin=True)
        with tempfile.TemporaryDirectory() as tmp:
            path = self._build_workbook(Path(tmp))
            call_command("import_payroll", path, "--commit", username="import_commit_admin")

        from payslip.models import Employee, PayrollRun, PayslipEntry
        employee = Employee.objects.get(name="Import Test One")
        self.assertEqual(employee.doj, datetime.date(2020, 1, 1))
        self.assertEqual(PayrollRun.objects.count(), 1)
        entry = PayslipEntry.objects.get()
        # The FIRST (original) row must win: lop_days=0, net=20000. A real
        # workbook has had two different rows for the same (period, name) -
        # the first one sits in the correctly-positioned block, the second
        # is a mislabeled duplicate - so first-row-wins is the correct rule.
        self.assertEqual(entry.lop_days, Decimal("0.00"))
        self.assertEqual(entry.net_payable, Decimal("20000.00"))

    def test_commit_is_safely_rerunnable(self):
        import tempfile
        from pathlib import Path
        from django.core.management import call_command

        make_member("import_rerun_admin", admin=True)
        with tempfile.TemporaryDirectory() as tmp:
            path = self._build_workbook(Path(tmp))
            call_command("import_payroll", path, "--commit", username="import_rerun_admin")
            call_command("import_payroll", path, "--commit", username="import_rerun_admin")

        from payslip.models import PayrollRun, PayslipEntry
        self.assertEqual(PayrollRun.objects.count(), 1)  # not duplicated
        self.assertEqual(PayslipEntry.objects.count(), 1)


# ===========================================================================
# Configurable salary engine
# ===========================================================================
class FormulaEngineTests(TestCase):
    """The safe AST evaluator: methods, functions, rounding, rejection of
    unsafe input, cycle detection and dependency ordering."""

    def test_fixed_percent_and_formula_methods(self):
        from decimal import Decimal
        from payslip.services.formula_engine import evaluate
        ctx = {"BASIC": Decimal("10000")}
        self.assertEqual(evaluate("5000", ctx), Decimal("5000"))
        self.assertEqual(evaluate("BASIC * 45 / 100", ctx), Decimal("4500"))
        self.assertEqual(evaluate("ROUND(BASIC * 0.333, 2)", ctx), Decimal("3330.00"))

    def test_functions(self):
        from decimal import Decimal
        from payslip.services.formula_engine import evaluate
        self.assertEqual(evaluate("ROUNDUP(100.1, 0)", {}), Decimal("101"))
        self.assertEqual(evaluate("ROUNDDOWN(100.9, 0)", {}), Decimal("100"))
        self.assertEqual(evaluate("MIN(3, 7, 5)", {}), Decimal("3"))
        self.assertEqual(evaluate("MAX(3, 7, 5)", {}), Decimal("7"))
        self.assertEqual(evaluate("IF(1, 10, 20)", {}), Decimal("10"))
        self.assertEqual(evaluate("IF(0, 10, 20)", {}), Decimal("20"))
        self.assertEqual(evaluate("IF(A > 5, 1, 0)", {"A": Decimal("6")}), Decimal("1"))

    def test_divide_by_zero_is_zero(self):
        from decimal import Decimal
        from payslip.services.formula_engine import evaluate
        self.assertEqual(evaluate("100 / TWD", {"TWD": Decimal("0")}), Decimal("0"))

    def test_rejects_unsafe_input(self):
        from payslip.services.formula_engine import FormulaError, validate_formula
        for bad in ('__import__("os")', "x.attr", "FOO(1)", "BASIC + UNKNOWN",
                    "[1, 2]", "lambda: 1"):
            with self.assertRaises(FormulaError, msg=bad):
                validate_formula(bad, {"BASIC"})

    def test_context_vars_are_legal_references(self):
        from payslip.services.formula_engine import validate_formula
        validate_formula("PACKAGE * 0.5", set())  # PACKAGE is a context var

    def test_circular_reference_detected(self):
        from payslip.services.formula_engine import FormulaError, order_components
        with self.assertRaises(FormulaError):
            order_components({"A": {"B"}, "B": {"A"}})

    def test_dependency_ordering(self):
        from payslip.services.formula_engine import order_components
        order = order_components({"A": {"B"}, "B": {"C"}, "C": set()})
        self.assertLess(order.index("C"), order.index("B"))
        self.assertLess(order.index("B"), order.index("A"))


class SalaryStructureEngineTests(TestCase):
    def setUp(self):
        from payslip.models import Organization, payroll_settings_for
        from payslip.services.structure_seed import seed_default_structure
        self.org = Organization.objects.create(company_name="Structure Test Org")
        self.settings = payroll_settings_for(self.org)
        self.structure = seed_default_structure(self.org)

    def test_matches_legacy_engine_field_for_field(self):
        """The linchpin: default structure == compute_entry across the
        verified worked examples."""
        from decimal import Decimal
        from payslip.services.payroll_calc import compute_entry
        from payslip.services.structure_calc import compute_entry_from_structure
        cases = [
            dict(monthly_package=Decimal("37000"), total_working_days=30,
                 is_pf_applicable=True),
            dict(monthly_package=Decimal("20000"), total_working_days=30,
                 salary_arrear_allowance=Decimal("5000"), is_esi_eligible=True,
                 is_pf_applicable=True),
            dict(monthly_package=Decimal("37000"), total_working_days=30,
                 lop_days=Decimal("2"), is_pf_applicable=True),
            dict(monthly_package=Decimal("15000"), total_working_days=31,
                 emp_leave_days=Decimal("1"), tds=Decimal("500"),
                 salary_advance=Decimal("1000"), is_esi_eligible=True),
        ]
        legacy_map = {"BASIC": "basic", "DA": "da", "HRA": "hra",
                      "TRANSPORT": "transport_allowance", "FOOD": "food_allowance",
                      "ESI_EMP": "esi_employee", "ESI_ER": "esi_employer",
                      "PF_EMP": "pf_employee", "PF_ER": "pf_employer"}
        for case in cases:
            old = compute_entry(settings=self.settings, **case)
            new = compute_entry_from_structure(self.structure, **case)
            amt = {c.code: c.amount for c in new.components}
            for code, field in legacy_map.items():
                self.assertEqual(amt.get(code, Decimal("0")), getattr(old, field),
                                 f"{code} for {case}")
            for field in ("present_days", "pay_days", "gross_salary",
                          "total_deductions", "net_payable"):
                self.assertEqual(getattr(new, field), getattr(old, field),
                                 f"{field} for {case}")

    def test_ctc_and_employer_totals(self):
        from decimal import Decimal
        from payslip.services.structure_calc import compute_entry_from_structure
        c = compute_entry_from_structure(
            self.structure, monthly_package=Decimal("37000"),
            total_working_days=30, is_pf_applicable=True)
        self.assertEqual(c.employer_contributions, Decimal("1800.00"))
        self.assertEqual(c.ctc, c.gross_salary + c.employer_contributions)
        self.assertEqual(c.net_payable, c.gross_salary - c.total_deductions)

    def test_custom_component_flows_into_gross(self):
        from decimal import Decimal
        from payslip.models import SalaryComponent
        from payslip.services.structure_calc import compute_entry_from_structure
        SalaryComponent.objects.create(
            structure=self.structure, code="SPECIAL", name="Special Allowance",
            kind=SalaryComponent.Kind.EARNING, calc_method=SalaryComponent.Method.FORMULA,
            formula="ROUND(BASIC * 0.10, 2)", rounding=SalaryComponent.Rounding.HALF_UP,
            decimals=2, include_in_gross=True, sequence=55)
        c = compute_entry_from_structure(
            self.structure, monthly_package=Decimal("37000"),
            total_working_days=30, is_pf_applicable=True)
        amt = {r.code: r.amount for r in c.components}
        self.assertEqual(amt["SPECIAL"], Decimal("1850.00"))  # 10% of 18500 basic
        self.assertEqual(c.gross_salary, Decimal("38850.00"))

    def test_versioning_selects_structure_by_period(self):
        import datetime
        from payslip.models import SalaryComponent, SalaryStructure, salary_structure_for
        v2 = SalaryStructure.objects.create(
            organization=self.org, label="2026 revision",
            effective_from=datetime.date(2026, 6, 1))
        SalaryComponent.objects.create(
            structure=v2, code="BASIC", name="Basic",
            kind=SalaryComponent.Kind.EARNING, calc_method=SalaryComponent.Method.FORMULA,
            formula="PACKAGE * 0.6 / TWD * PAY_DAYS",
            rounding=SalaryComponent.Rounding.HALF_UP, decimals=0, sequence=10)
        self.assertEqual(
            salary_structure_for(self.org, datetime.date(2026, 5, 1)).pk,
            self.structure.pk)
        self.assertEqual(
            salary_structure_for(self.org, datetime.date(2026, 6, 1)).pk, v2.pk)

    def test_default_structure_autoseeds_once(self):
        from payslip.models import SalaryStructure, salary_structure_for
        again = salary_structure_for(self.org)
        self.assertEqual(again.pk, self.structure.pk)
        self.assertEqual(
            SalaryStructure.objects.filter(organization=self.org).count(), 1)


class SalaryStructureViewTests(TestCase):
    def setUp(self):
        self.admin = make_member("struct_admin", admin=True)
        self.org = self.admin.membership.organization
        self.client.force_login(self.admin)

    def test_structure_page_seeds_and_renders(self):
        resp = self.client.get(reverse("salary_structure"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "BASIC")
        self.assertContains(resp, "PF_EMP")

    def test_member_without_admin_cannot_open_structure(self):
        member = make_member("struct_member", org=self.org, payroll=True)
        self.client.force_login(member)
        self.assertEqual(self.client.get(reverse("salary_structure")).status_code, 403)

    def test_add_component_via_form(self):
        from payslip.models import salary_structure_for
        salary_structure_for(self.org)
        resp = self.client.post(reverse("salary_structure"), {
            "action": "save_component", "code": "SPECIAL", "name": "Special",
            "kind": "EARNING", "calc_method": "FORMULA", "amount": "0",
            "percent": "0", "base_code": "", "formula": "ROUND(BASIC * 0.05, 2)",
            "rounding": "HALF_UP", "decimals": "2", "include_in_gross": "on",
            "statutory_type": "NONE", "sequence": "55", "is_active": "on",
        })
        self.assertEqual(resp.status_code, 302)
        struct = salary_structure_for(self.org)
        self.assertTrue(struct.components.filter(code="SPECIAL").exists())

    def test_self_referential_formula_rejected_by_form(self):
        from payslip.models import salary_structure_for
        salary_structure_for(self.org)
        self.client.post(reverse("salary_structure"), {
            "action": "save_component", "code": "LOOP", "name": "Loop",
            "kind": "EARNING", "calc_method": "FORMULA", "amount": "0",
            "percent": "0", "base_code": "", "formula": "LOOP + 1",
            "rounding": "HALF_UP", "decimals": "2",
            "statutory_type": "NONE", "sequence": "60", "is_active": "on",
        })
        struct = salary_structure_for(self.org)
        self.assertFalse(struct.components.filter(code="LOOP").exists())


class SalaryEngineBackwardCompatTests(TestCase):
    def setUp(self):
        from decimal import Decimal
        from payslip.models import Employee
        self.admin = make_member("bc_admin", admin=True, payroll=True)
        self.org = self.admin.membership.organization
        self.client.force_login(self.admin)
        self.employee = Employee.objects.create(
            organization=self.org, employee_code="EMP-0001", name="BC Test",
            current_monthly_package=Decimal("30000"), is_pf_applicable=True)

    def test_new_run_populates_components_and_ctc(self):
        from payslip.models import PayslipEntry
        self.client.post(reverse("payroll_run_create"), {"period": "2026-09"})
        entry = PayslipEntry.objects.get()
        self.assertGreater(entry.basic, 0)
        self.assertGreater(entry.gross_salary, 0)
        self.assertGreater(entry.ctc, entry.gross_salary)  # employer PF added
        self.assertTrue(entry.component_amounts.filter(code="PF_ER").exists())

    def test_finalized_run_is_never_recomputed(self):
        from decimal import Decimal
        from payslip.models import PayrollRun, PayslipEntry
        self.client.post(reverse("payroll_run_create"), {"period": "2026-09"})
        run = PayrollRun.objects.get()
        entry = PayslipEntry.objects.get()
        PayslipEntry.objects.filter(pk=entry.pk).update(net_payable=Decimal("99999"))
        PayrollRun.objects.filter(pk=run.pk).update(status=PayrollRun.Status.FINALIZED)
        self.client.get(reverse("payroll_run_detail", args=[run.pk]))
        entry.refresh_from_db()
        self.assertEqual(entry.net_payable, Decimal("99999"))


class PayrollRecalculateTests(TestCase):
    def setUp(self):
        from decimal import Decimal
        from payslip.models import Employee
        self.admin = make_member("recalc_admin", admin=True, payroll=True)
        self.org = self.admin.membership.organization
        self.client.force_login(self.admin)
        self.e1 = Employee.objects.create(
            organization=self.org, employee_code="EMP-0001", name="Alpha",
            current_monthly_package=Decimal("30000"), is_pf_applicable=True)
        self.e2 = Employee.objects.create(
            organization=self.org, employee_code="EMP-0002", name="Beta",
            current_monthly_package=Decimal("20000"))
        self.client.post(reverse("payroll_run_create"), {"period": "2026-09"})
        from payslip.models import PayrollRun
        self.run = PayrollRun.objects.get()

    def test_recalc_removes_deactivated_employee(self):
        from payslip.models import PayslipEntry
        self.assertEqual(self.run.entries.count(), 2)
        self.e2.is_active = False
        self.e2.save(update_fields=["is_active"])
        self.client.post(reverse("payroll_run_recalculate", args=[self.run.pk]))
        self.assertEqual(self.run.entries.count(), 1)
        self.assertFalse(self.run.entries.filter(employee=self.e2).exists())

    def test_recalc_adds_new_active_employee(self):
        from decimal import Decimal
        from payslip.models import Employee
        Employee.objects.create(
            organization=self.org, employee_code="EMP-0003", name="Gamma",
            current_monthly_package=Decimal("25000"))
        self.client.post(reverse("payroll_run_recalculate", args=[self.run.pk]))
        self.assertEqual(self.run.entries.count(), 3)
        added = self.run.entries.get(employee__name="Gamma")
        self.assertGreater(added.basic, 0)
        self.assertTrue(added.component_amounts.exists())

    def test_recalc_refreshes_package_but_keeps_manual_attendance(self):
        from decimal import Decimal
        from payslip.models import PayslipEntry
        entry = self.run.entries.get(employee=self.e1)
        # Enter a manual LOP on the row, then raise the master package.
        PayslipEntry.objects.filter(pk=entry.pk).update(lop_days=Decimal("3"),
                                                        total_working_days=30)
        self.e1.current_monthly_package = Decimal("36000")
        self.e1.save(update_fields=["current_monthly_package"])
        self.client.post(reverse("payroll_run_recalculate", args=[self.run.pk]))
        entry.refresh_from_db()
        self.assertEqual(entry.monthly_package, Decimal("36000"))  # refreshed
        self.assertEqual(entry.lop_days, Decimal("3"))             # preserved
        # Basic recomputed on the new package, prorated for 3 LOP of 30 days.
        self.assertEqual(entry.pay_days, Decimal("27"))
        self.assertGreater(entry.ctc, entry.gross_salary)

    def test_recalc_noop_on_finalized_run(self):
        from decimal import Decimal
        from payslip.models import PayrollRun, PayslipEntry
        PayslipEntry.objects.filter(run=self.run).update(net_payable=Decimal("55555"))
        PayrollRun.objects.filter(pk=self.run.pk).update(
            status=PayrollRun.Status.FINALIZED)
        resp = self.client.post(reverse("payroll_run_recalculate", args=[self.run.pk]))
        self.assertEqual(resp.status_code, 302)
        for e in self.run.entries.all():
            self.assertEqual(e.net_payable, Decimal("55555"))  # untouched

    def test_recalc_requires_admin(self):
        member = make_member("recalc_member", org=self.org, payroll=True)
        self.client.force_login(member)
        resp = self.client.post(reverse("payroll_run_recalculate", args=[self.run.pk]))
        self.assertEqual(resp.status_code, 403)
